"""Reporte Excel: Resumen · Débito · Crédito · Cuotas · Prepago · Pendientes · Conciliados · Parámetros."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..motor import categorias as cat
from .datos import armar

TINTA = "1B2430"
MILES = "#,##0"
FECHA = "DD/MM/YYYY"
BORDE = Border(bottom=Side(style="thin", color="D6DCE3"))
ENCABEZADO_FONDO = PatternFill("solid", fgColor=TINTA)
ENCABEZADO_FUENTE = Font(bold=True, color="FFFFFF")
TIPO_COLOR = {"Operación": "B42318", "Dato": "6B3FA0", "Pendiente": "475467", "Aviso": "5B6675", "OK": "2F7D4F"}

COLUMNAS_VENTA = ["Fecha venta", "Código", "Local", "Monto Transbank", "Monto bordero", "Diferencia", "Cuota",
                  "Categoría", "Tipo", "Código sugerido", "Mensaje", "Nota", "Revisado", "Observación"]


def _hex(color: str) -> str:
    return color.lstrip("#").upper()


class Hoja:
    """Ayuda para escribir una hoja fila por fila con formato consistente."""

    def __init__(self, libro, titulo, color=None):
        self.ws = libro.create_sheet(titulo)
        self.fila = 1
        if color:
            self.ws.sheet_properties.tabColor = _hex(color)

    def texto(self, valor, negrita=False, tamano=11, color=None):
        celda = self.ws.cell(self.fila, 1, valor)
        celda.font = Font(bold=negrita, size=tamano, color=_hex(color) if color else None)
        self.fila += 1

    def salto(self, n=1):
        self.fila += n

    def tabla(self, columnas, filas, montos=(), fechas=(), anchos=None, filtro=False, congelar=False,
              color_tipo_col=None):
        inicio = self.fila
        for c, titulo in enumerate(columnas, 1):
            celda = self.ws.cell(self.fila, c, titulo)
            celda.fill, celda.font = ENCABEZADO_FONDO, ENCABEZADO_FUENTE
            celda.alignment = Alignment(vertical="center", wrap_text=True)
        self.fila += 1
        for valores in filas:
            for c, valor in enumerate(valores, 1):
                celda = self.ws.cell(self.fila, c, valor)
                celda.border = BORDE
                titulo = columnas[c - 1]
                if titulo in montos:
                    celda.number_format = MILES
                elif titulo in fechas:
                    celda.number_format = FECHA
                if titulo in ("Mensaje", "Observación", "Corregir"):
                    celda.alignment = Alignment(wrap_text=True, vertical="top")
                if color_tipo_col and titulo == color_tipo_col and valor in TIPO_COLOR:
                    celda.font = Font(bold=True, color=TIPO_COLOR[valor])
            self.fila += 1
        fin = self.fila - 1
        if filtro and fin > inicio:
            self.ws.auto_filter.ref = f"A{inicio}:{get_column_letter(len(columnas))}{fin}"
        if congelar:
            self.ws.freeze_panes = f"A{inicio + 1}"
        for c, ancho in enumerate(anchos or [], 1):
            self.ws.column_dimensions[get_column_letter(c)].width = ancho
        self.fila += 1


def _encabezado(h: Hoja, d: dict, titulo: str, color=None):
    e = d["encabezado"]
    h.texto(e["empresa"], negrita=True)
    h.texto(f"{e['departamento']} · RUT {e['rut']}")
    h.texto(titulo, negrita=True, tamano=14, color=color)
    h.texto(f"{e['subtitulo']} · {e['cuenta']} · Generado {e['generado']:%d/%m/%Y %H:%M}")
    h.salto()


def _fila_venta(v) -> list:
    return [v.fecha_venta, v.codigo_autorizacion, v.local, v.monto_original, v.monto_bordero, v.diferencia,
            v.cuota, cat.ETIQUETA.get(v.categoria, v.categoria), v.tipo, v.codigo_sugerido, v.mensaje, v.nota,
            "Sí" if v.revisado else "", v.observacion]


def _resumen(libro, d):
    h = Hoja(libro, "Resumen")
    _encabezado(h, d, d["encabezado"]["titulo"])
    columnas = ["Medio de pago", "Ventas Transbank", "Comisión + IVA", "Abono Transbank", "Abono en cartola",
                "Diferencial", "N° ventas", "Ventas con falla"]
    filas = [[r["etiqueta"], r.get("ventas", 0), r.get("comision_iva", 0), r.get("abono_tbk", 0),
              r.get("abono_cartola", 0), r.get("diferencial", 0), r.get("n_ventas", 0), r.get("n_falla", 0)]
             for r in d["resumen_medios"]]
    t = d["total"]
    filas.append(["Total", t.get("ventas", 0), t.get("comision_iva", 0), t.get("abono_tbk", 0),
                  t.get("abono_cartola", 0), t.get("diferencial", 0), t.get("n_ventas", 0), t.get("n_falla", 0)])
    h.tabla(columnas, filas, montos=columnas[1:6], anchos=[34, 18, 16, 18, 18, 14, 12, 16])
    for i, r in enumerate(d["resumen_medios"]):
        h.ws.cell(h.fila - 2 - len(d["resumen_medios"]) + i, 1).font = Font(bold=True, color=_hex(r["color"]))
    h.ws.cell(h.fila - 2, 1).font = Font(bold=True)

    h.texto("Matriz de resultados (cantidad de registros)", negrita=True, tamano=12)
    columnas = ["Categoría", "Tipo"] + [cat.ETIQUETA_MEDIO[m] for m in cat.MEDIOS] + ["Total", "Monto total"]
    filas = [[f["etiqueta"], f["tipo"], *[c["n"] for c in f["celdas"]], f["total_n"], f["total_monto"]]
             for f in d["matriz"]]
    h.tabla(columnas, filas, montos=["Monto total"], color_tipo_col="Tipo")
    h.texto("Red Compra (RCO) puede ser débito o prepago; sus pendientes se muestran bajo Débito.")
    h.texto("Cuadratura correcta: en cada medio de pago la suma de las categorías es igual a Ventas Transbank."
            if d["cuadra"] else "ATENCIÓN: la suma de las categorías no coincide con Ventas Transbank.",
            negrita=not d["cuadra"])
    if d["cartola_sin_archivo"]:
        h.salto()
        h.texto(f"Abonos de la cartola sin archivo Transbank ({len(d['cartola_sin_archivo'])}): cargar los archivos "
                "'Abonos por día' de esas fechas.", negrita=True, color="#B42318")
        h.tabla(["Fecha", "Descripción en la cartola", "Medio probable", "Monto"],
                [[m.fecha, m.descripcion, m.medio_probable, m.abono] for m in d["cartola_sin_archivo"]],
                montos=["Monto"], fechas=["Fecha"])


def _capitulo(libro, d, medio):
    cap = d["capitulos"][medio]
    h = Hoja(libro, cap["etiqueta"], cap["color"])
    _encabezado(h, d, f"{cap['etiqueta']} — {d['conciliacion'].etiqueta}", cap["color"])
    r = cap["resumen"]
    h.texto(f"Ventas: {r.get('n_ventas', 0)} · Con falla: {r.get('n_falla', 0)} · "
            f"Ventas Transbank ${r.get('ventas', 0):,} · Abono ${r.get('abono_tbk', 0):,}".replace(",", "."))
    for a in cap["abonos"]:
        estado = "cuadra con la cartola" if a.categoria == cat.OK else "NO cuadra con la cartola"
        h.texto(f"Abono del {a.fecha_abono:%d/%m/%Y}: ${a.total_abono:,} — {estado}. {a.mensaje}".replace(",", "."))
    h.salto()
    filas = [_fila_venta(v) for v in cap["fallas"] + cap["avisos"]]
    h.tabla(COLUMNAS_VENTA, filas, montos=["Monto Transbank", "Monto bordero", "Diferencia"],
            fechas=["Fecha venta"], filtro=True, congelar=True, color_tipo_col="Tipo",
            anchos=[12, 12, 24, 15, 15, 12, 7, 26, 11, 15, 70, 26, 9, 30])
    if cap["otros"]:
        h.texto("Otros movimientos (no son ventas)", negrita=True)
        h.tabla(["Fecha venta", "Tipo de movimiento", "Código", "Monto"],
                [[v.fecha_venta, v.tipo_movimiento, v.codigo_autorizacion, v.monto_original] for v in cap["otros"]],
                montos=["Monto"], fechas=["Fecha venta"])


def _cuotas(libro, d):
    h = Hoja(libro, "Cuotas", cat.COLOR_MEDIO[cat.CREDITO])
    _encabezado(h, d, f"Cuotas — {d['conciliacion'].etiqueta}", cat.COLOR_MEDIO[cat.CREDITO])
    columnas = ["Fecha venta", "Código", "Medio", "Local", "Cuota", "Monto venta", "Monto de esta cuota",
                "Categoría", "Mensaje"]
    filas = [[v.fecha_venta, v.codigo_autorizacion, cat.ETIQUETA_MEDIO[v.medio], v.local, v.cuota,
              v.monto_original, v.monto_abono_valido, cat.ETIQUETA.get(v.categoria), v.mensaje] for v in d["cuotas"]]
    h.tabla(columnas, filas, montos=["Monto venta", "Monto de esta cuota"], fechas=["Fecha venta"],
            anchos=[12, 12, 10, 24, 7, 15, 18, 26, 70])


def _pendientes(libro, d):
    h = Hoja(libro, "Pendientes")
    _encabezado(h, d, f"Pendientes abiertos al cierre — {d['conciliacion'].etiqueta}")
    h.texto(f"Pagos con tarjeta del bordero sin abono de Transbank. Cerrados este mes: {d['n_pendientes_cerrados']}.")
    h.salto()
    columnas = ["Fecha", "Cod. Deb.", "Medio", "Documento", "Monto", "Hab", "Cuenta", "Mes de origen", "Mensaje",
                "Observación", "Corregir"]
    filas = [[p.pago.fecha, p.pago.cod_deb, cat.ETIQUETA_MEDIO[p.medio], p.pago.documento, abs(p.pago.valor),
              p.pago.hab, p.pago.cuenta, p.origen.etiqueta, p.mensaje, p.observacion, ""] for p in d["pendientes"]]
    h.tabla(columnas, filas, montos=["Monto"], fechas=["Fecha"], filtro=True, congelar=True,
            anchos=[12, 10, 10, 14, 13, 8, 12, 16, 70, 30, 20])
    if d["invalidos_bordero"]:
        h.texto("Datos inválidos del bordero (pagos con tarjeta con monto positivo / reversas)", negrita=True)
        h.tabla(["Fecha", "Cod. Deb.", "Documento", "Monto", "Descripción"],
                [[p.fecha, p.cod_deb, p.documento, p.valor, p.descripcion] for p in d["invalidos_bordero"]],
                montos=["Monto"], fechas=["Fecha"])


def _conciliados(libro, d):
    h = Hoja(libro, "Conciliados")
    _encabezado(h, d, f"Ventas conciliadas — {d['conciliacion'].etiqueta}")
    columnas = ["Medio de pago", "Fecha venta", "Código", "Local", "Monto Transbank", "Monto bordero",
                "Diferencia", "Documento bordero", "Nota"]
    filas = []
    for m in cat.MEDIOS:
        for v in d["capitulos"][m]["conciliados"]:
            filas.append([cat.ETIQUETA_MEDIO[m], v.fecha_venta, v.codigo_autorizacion, v.local, v.monto_original,
                          v.monto_bordero, v.diferencia, v.pago.documento if v.pago else "", v.nota])
    h.tabla(columnas, filas, montos=["Monto Transbank", "Monto bordero", "Diferencia"], fechas=["Fecha venta"],
            filtro=True, congelar=True, anchos=[14, 12, 12, 24, 15, 15, 12, 18, 40])


def _parametros(libro, d):
    h = Hoja(libro, "Parámetros")
    _encabezado(h, d, "Parámetros y archivos procesados")
    p = d["parametros"]
    filas = [
        ["Tolerancia en pesos", p.get("tolerancia_clp")],
        ["Similitud mínima de códigos", p.get("umbral_similitud")],
        ["Días hábiles de holgura en cartola", p.get("dias_holgura")],
        ["Días máximos entre venta y bordero", p.get("ventana_dias")],
        ["Textos de abonos en cartola", ", ".join(p.get("patrones_cartola", []))],
        ["Medio probable según cartola", "; ".join(f"{t} = {m}" for t, m in p.get("reglas_medio_cartola", []))],
        ["Omitir archivos en moneda extranjera", "Sí" if p.get("omitir_moneda_extranjera") else "No"],
    ] + [[f"Equivalencia {k}", ", ".join(cat.ETIQUETA_MEDIO.get(m, m) for m in v)]
         for k, v in p.get("equivalencias", {}).items()]
    h.tabla(["Parámetro", "Valor"], filas, anchos=[36, 40])
    h.texto("Archivos", negrita=True)
    h.tabla(["Archivo", "Tipo", "Medio", "Desde", "Hasta", "Fecha abono", "Filas", "Total abono", "SHA-256"],
            [[a.nombre, a.get_tipo_display() + (f" (omitido, {a.moneda})" if a.moneda_extranjera else ""),
              a.etiqueta_medio, a.desde, a.hasta, a.fecha_abono, a.filas, a.total_abono, a.sha256[:16]]
             for a in d["archivos"]],
            montos=["Total abono"], fechas=["Desde", "Hasta", "Fecha abono"])


def generar(conciliacion) -> bytes:
    d = armar(conciliacion)
    libro = Workbook()
    libro.remove(libro.active)
    _resumen(libro, d)
    _capitulo(libro, d, cat.DEBITO)
    _capitulo(libro, d, cat.CREDITO)
    _cuotas(libro, d)
    _capitulo(libro, d, cat.PREPAGO)
    _pendientes(libro, d)
    _conciliados(libro, d)
    _parametros(libro, d)
    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()
