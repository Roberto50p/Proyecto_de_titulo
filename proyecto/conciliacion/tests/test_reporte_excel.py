from io import BytesIO

from openpyxl import load_workbook

from conciliacion.reportes.excel import generar

HOJAS = ["Resumen", "Débito", "Crédito", "Cuotas", "Prepago", "Pendientes", "Conciliados", "Parámetros"]


def filas_tabla(hoja, encabezado_inicial):
    """Filas de datos de la tabla cuyo encabezado empieza con `encabezado_inicial`."""
    filas = list(hoja.iter_rows(values_only=True))
    inicio = next(i for i, f in enumerate(filas) if f and f[0] == encabezado_inicial)
    titulos = filas[inicio]
    datos = []
    for f in filas[inicio + 1:]:
        if f[0] is None:
            break
        datos.append(dict(zip(titulos, f)))
    return datos


def test_excel_septiembre(septiembre_procesado):
    libro = load_workbook(BytesIO(generar(septiembre_procesado)))
    assert libro.sheetnames == HOJAS

    resumen = filas_tabla(libro["Resumen"], "Medio de pago")
    por_medio = {f["Medio de pago"]: f for f in resumen}
    assert por_medio["Total"]["Abono Transbank"] == 5_269_462 + 2_597_371 + 861_527
    assert por_medio["Débito"]["Diferencial"] == 0
    assert por_medio["Crédito"]["Ventas con falla"] == 5

    debito = filas_tabla(libro["Débito"], "Fecha venta")
    fallas = [f for f in debito if f["Tipo"] in ("Operación", "Dato")]
    assert len(fallas) == 8
    assert any(f["Código"] == "F0FUAZ" and f["Código sugerido"] == "0FOFUAZ" for f in fallas)
    assert libro["Débito"].auto_filter.ref is not None
    assert libro["Débito"].freeze_panes is not None

    cuotas = filas_tabla(libro["Cuotas"], "Fecha venta")
    assert {f["Cuota"] for f in cuotas} == {"1/3", "3/3"}

    pendientes = filas_tabla(libro["Pendientes"], "Fecha")
    assert len(pendientes) == 123
    assert "Observación" in pendientes[0] and "Corregir" in pendientes[0]

    conciliados = filas_tabla(libro["Conciliados"], "Medio de pago")
    assert len(conciliados) == 97
