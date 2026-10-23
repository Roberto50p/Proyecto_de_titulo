from conciliacion.models import PagoBordero, Pendiente
from conciliacion.servicio import procesar


def test_bordero_traslapado_no_duplica(crear_conciliacion):
    procesar(crear_conciliacion(2026, 8, ["BORDERO 04-09 A 07-09.xlsx"]))
    total = PagoBordero.objects.count()
    procesar(crear_conciliacion(2026, 9, ["BORDERO 04-09 A 07-09.xlsx"]))
    assert PagoBordero.objects.count() == total == 228


def test_pendiente_se_cierra_mes_siguiente(crear_conciliacion):
    agosto = procesar(crear_conciliacion(2026, 8, ["BORDERO 04-09 A 07-09.xlsx"]))
    assert agosto.pendientes.filter(abierto=True).count() == 227

    septiembre = procesar(crear_conciliacion())
    venta = septiembre.ventas.get(codigo_autorizacion="003098")
    assert venta.categoria == "OK" and "cierra pendiente de Agosto 2026" in venta.nota
    assert Pendiente.objects.get(pago=venta.pago).cierre == septiembre
    assert septiembre.resumen["n_pendientes_cerrados"] == 104
    # En septiembre no se reabren como pendientes nuevos: ya estaban abiertos desde agosto.
    assert septiembre.pendientes.count() == 0
    assert Pendiente.objects.filter(abierto=True).count() == 123


def test_reprocesar_conserva_observacion(septiembre_procesado):
    conc = septiembre_procesado
    v = conc.ventas.get(codigo_autorizacion="007626")
    v.observacion, v.revisado = "Se cobró propina aparte", True
    v.save()
    procesar(conc)
    v = conc.ventas.get(codigo_autorizacion="007626")
    assert (v.observacion, v.revisado) == ("Se cobró propina aparte", True)
    assert conc.ventas.count() == 141
    assert PagoBordero.objects.filter(conciliacion=conc).count() == 104


# ---- Correcciones de la revisión final

def test_reprocesar_mes_anterior_conserva_pendientes_cerrados(crear_conciliacion):
    agosto = procesar(crear_conciliacion(2026, 8, ["BORDERO 04-09 A 07-09.xlsx"]))
    septiembre = procesar(crear_conciliacion())
    cerrado = Pendiente.objects.get(pago__documento_norm="3098")
    cerrado.observacion = "Revisado con Transbank"
    cerrado.save()
    procesar(agosto)
    cerrado = Pendiente.objects.get(pago__documento_norm="3098")
    assert cerrado.cierre == septiembre and cerrado.observacion == "Revisado con Transbank"
    assert Pendiente.objects.filter(cierre=septiembre).count() == 104


def test_reporte_de_un_mes_no_cambia_al_procesar_el_siguiente(crear_conciliacion):
    from conciliacion.reportes.datos import armar

    agosto = procesar(crear_conciliacion(2026, 8, ["BORDERO 04-09 A 07-09.xlsx"]))
    assert len(armar(agosto)["pendientes"]) == 227
    procesar(crear_conciliacion())
    assert len(armar(agosto)["pendientes"]) == 227


def test_abonos_de_cartola_sin_archivo(septiembre_procesado):
    sin_archivo = septiembre_procesado.resumen["cartola_sin_archivo"]
    assert sin_archivo["n"] == 16 and sin_archivo["monto"] == 26_036_232


def test_matriz_cuadra_con_ventas_transbank(septiembre_procesado):
    from conciliacion.motor import categorias as cat

    r = septiembre_procesado.resumen
    ventas = {k for k in cat.ORDEN_CATEGORIAS if k not in (cat.SIN_ABONO, cat.ABONO_CARTOLA, cat.DATO_INVALIDO)}
    for medio in cat.MEDIOS:
        suma = sum(r["matriz"][k][medio]["monto"] for k in ventas)
        assert suma == r["medios"][medio]["ventas"]


def test_documento_corregido_en_bordero_no_crea_pendiente_falso(septiembre_procesado, crear_conciliacion,
                                                                datos_reales, tmp_path):
    from django.core.files import File
    from openpyxl import load_workbook

    from conciliacion.models import ArchivoCargado
    from conciliacion.servicio import leer_archivo, sha256_de

    libro = load_workbook(datos_reales / "BORDERO 04-09 A 07-09.xlsx")
    hoja = libro.active
    for fila in hoja.iter_rows():
        for celda in fila:
            if celda.value == "0FOFUAZ":
                celda.value = "F0FUAZ"
    corregido = tmp_path / "bordero corregido.xlsx"
    libro.save(corregido)

    octubre = crear_conciliacion(2026, 10, [])
    with open(corregido, "rb") as f:
        archivo = File(f, name=corregido.name)
        ac = ArchivoCargado(conciliacion=octubre, nombre=corregido.name, sha256=sha256_de(archivo))
        ac.archivo.save(corregido.name, archivo, save=False)
        ac.save()
    leer_archivo(ac)
    procesar(octubre)
    assert PagoBordero.objects.count() == 228
    pago = PagoBordero.objects.get(documento="F0FUAZ")
    assert pago.conciliacion == septiembre_procesado
    assert not octubre.pendientes.exists()
